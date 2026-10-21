"""Generate a blank, auditable Douyin creator-BD workbook; no company data is embedded."""
from __future__ import annotations

from datetime import datetime
from decimal import Decimal, ROUND_CEILING, ROUND_HALF_UP
from pathlib import Path
import json
import re
import xlsxwriter

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "达人带货商务台账.xlsx"
ROWS = 300
LAST = ROWS + 1


def calculate_quote(*, net_income, commission_base, product_cost, packaging, shipping,
                    other_cost, commission_rate, platform_rate, platform_base,
                    fixed_fee, sample_budget, expected_orders=None):
    """Independent reference for the sheet's estimate; not an Excel formula evaluator."""
    raw = [net_income, commission_base, product_cost, packaging, shipping, other_cost,
           commission_rate, platform_rate, platform_base, fixed_fee, sample_budget]
    if any(value is None or value == "" for value in raw):
        return {"state": "待补充"}
    values = list(map(lambda x: Decimal(str(x)), raw))
    if min(values) < 0 or values[6] > 1 or values[7] > 1:
        return {"state": "检查输入"}
    income, base, goods, pack, freight, other, commission, fee_rate, fee_base, fixed, samples = values
    contribution = (income - goods - pack - freight - other - base * commission - fee_base * fee_rate).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
    upfront = fixed + samples
    if contribution <= 0:
        return {"state": "不可保本", "unit_contribution": contribution, "upfront": upfront}
    break_even = int((upfront / contribution).to_integral_value(rounding=ROUND_CEILING))
    result = {"state": "可估算", "unit_contribution": contribution, "upfront": upfront, "break_even_orders": break_even}
    if expected_orders is not None:
        orders = Decimal(str(expected_orders))
        if orders < 0 or orders != orders.to_integral_value():
            return {"state": "检查输入"}
        result["scenario_contribution"] = orders * contribution - upfront
    return result


def make_workbook(output=OUTPUT):
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = xlsxwriter.Workbook(output, {"strings_to_urls": False})
    wb.set_properties({"title": "达人短视频带货商务台账", "subject": "食品类目 BD 试运行模板", "author": "BD 工作台模板", "comments": "空白模板；示例与真实经营数据分开。"})
    wb.set_calc_mode("auto")
    base = {"font_name": "微软雅黑", "font_size": 10, "valign": "top", "text_wrap": True}
    fmt = {
        "header": wb.add_format({**base, "bold": True, "font_color": "#FFFFFF", "bg_color": "#243E3D", "valign": "vcenter", "border": 1, "border_color": "#49605C"}),
        "input": wb.add_format({**base, "font_color": "#285B79", "bg_color": "#F3F8FC", "border": 1, "border_color": "#E2EBED"}),
        "text": wb.add_format({**base, "border": 1, "border_color": "#E2EBED"}),
        "formula": wb.add_format({**base, "font_color": "#83632E", "bg_color": "#FFF8E8", "border": 1, "border_color": "#EEE5D2"}),
        "date": wb.add_format({**base, "font_color": "#285B79", "bg_color": "#F3F8FC", "border": 1, "border_color": "#E2EBED", "num_format": "yyyy-mm-dd"}),
        "money": wb.add_format({**base, "font_color": "#285B79", "bg_color": "#F3F8FC", "border": 1, "border_color": "#E2EBED", "num_format": '#,##0.00;[Red]-#,##0.00'}),
        "percent": wb.add_format({**base, "font_color": "#285B79", "bg_color": "#F3F8FC", "border": 1, "border_color": "#E2EBED", "num_format": "0.00%"}),
        "integer": wb.add_format({**base, "font_color": "#285B79", "bg_color": "#F3F8FC", "border": 1, "border_color": "#E2EBED", "num_format": "0"}),
        "formula_money": wb.add_format({**base, "font_color": "#83632E", "bg_color": "#FFF8E8", "border": 1, "border_color": "#EEE5D2", "num_format": '#,##0.00;[Red]-#,##0.00'}),
        "bad": wb.add_format({"bg_color": "#FCE8E4", "font_color": "#A83C2C"}),
        "good": wb.add_format({"bg_color": "#E4F0E8", "font_color": "#27613B"}),
        "muted": wb.add_format({**base, "font_color": "#737D83", "bg_color": "#F4F5F6"}),
        "title": wb.add_format({**base, "font_size": 19, "bold": True, "font_color": "#243E3D"}),
        "section": wb.add_format({**base, "font_size": 12, "bold": True, "font_color": "#243E3D", "bg_color": "#E9F0EB"}),
        "warning": wb.add_format({**base, "bold": True, "font_color": "#9B5C24", "bg_color": "#FFF3DA"}),
    }

    def setup(name, headers, widths=None, freeze=2):
        ws = wb.add_worksheet(name)
        ws.hide_gridlines(2)
        ws.set_tab_color("#789F8D")
        ws.set_default_row(34)
        ws.set_row(0, 48)
        ws.freeze_panes(1, freeze)
        ws.set_landscape()
        ws.fit_to_pages(1, 0)
        ws.repeat_rows(0)
        ws.set_footer("&L达人带货商务 · 空白模板&R第 &P / &N 页")
        ws.set_column(0, len(headers) - 1, 18, fmt["input"])
        if widths:
            for col, width in widths.items():
                ws.set_column(col, col, width, fmt["input"])
        for col, text in enumerate(headers): ws.write(0, col, text, fmt["header"])
        ws.autofilter(0, 0, ROWS, len(headers) - 1)
        return ws

    def dates(ws, columns):
        for col in columns:
            ws.set_column(col, col, 14, fmt["date"])
            ws.data_validation(1, col, ROWS, col, {"validate": "date", "criteria": "between", "minimum": datetime(2000, 1, 1), "maximum": datetime(2100, 12, 31), "ignore_blank": True, "input_title": "真实业务日期", "input_message": "按 yyyy-mm-dd 填写；不知道先留空。", "error_type": "stop", "error_title": "日期不正确", "error_message": "请填写 2000–2100 年间有效日期。"})

    def numbers(ws, columns, *, percent=False, integer=False):
        for col in columns:
            cell_fmt = fmt["percent"] if percent else fmt["integer"] if integer else fmt["money"]
            ws.set_column(col, col, 17, cell_fmt)
            ws.data_validation(1, col, ROWS, col, {"validate": "integer" if integer else "decimal", "criteria": "between", "minimum": 0, "maximum": 1 if percent else 1_000_000_000, "ignore_blank": True, "error_type": "stop", "error_title": "输入范围不正确", "error_message": "百分比请输入 0%–100%。" if percent else "请填写非负数；确认没有费用时填 0，未知时留空。"})

    def dropdown(ws, col, name):
        ws.data_validation(1, col, ROWS, col, {"validate": "list", "source": f"={name}", "ignore_blank": True, "error_type": "stop", "error_title": "请选择有效状态", "error_message": "从下拉选项选择；参数字典中维护选项。"})

    def formula(ws, row, col, expression, money=False):
        ws.write_formula(row - 1, col, expression, fmt["formula_money"] if money else fmt["formula"], "")

    # Dictionary and food/product evidence.
    parameter_headers = ["分类", "资料项", "当前已核实值（未知留空）", "单位 / 适用口径", "证据或资料位置", "核对负责人", "核对日期", "有效 / 到期日期", "状态", "下次核对日期", "备注 / 待确认问题"]
    params = setup("参数字典", parameter_headers, {0: 14, 1: 28, 2: 32, 3: 32, 4: 30, 10: 38}, freeze=2)
    dates(params, [6, 7, 9])
    lists = {
        "LeadStatus": ["待筛选", "待触达", "已触达", "沟通中", "待审批", "待寄样", "试用中", "待发布", "已发布观察", "待结算", "复盘完成", "终止"],
        "CollabType": ["纯佣", "固定费+佣金", "固定费", "其他待核实"],
        "ApprovalStatus": ["待申请", "审批中", "已批准", "已拒绝", "不涉及"],
        "SampleStatus": ["待审批", "待发出", "运输中", "已签收", "待反馈", "已完成", "异常", "已取消"],
        "ReviewStatus": ["待发布", "已发布观察", "待结算", "已复盘", "异常", "终止"],
        "YesNo": ["是", "否"],
        "EvidenceStatus": ["待核实", "已确认", "已过期", "不适用"],
    }
    for offset, (name, values) in enumerate(lists.items(), start=12):
        params.write(0, offset, name, fmt["header"])
        for i, val in enumerate(values, start=1): params.write(i, offset, val)
        letter = xlsxwriter.utility.xl_col_to_name(offset)
        wb.define_name(name, f"='参数字典'!${letter}$2:${letter}${len(values)+1}")
    params.set_column(12, 18, None, None, {"hidden": True})
    dropdown(params, 8, "EvidenceStatus")
    evidence = [
        ("商品", "SKU 与商品全名", "区分规格、口味、组合、版本，防止错链"),
        ("商品", "规格 / 净含量 / 件数", "与实物包装、商品详情页一致"),
        ("商品", "配料表", "用当前版本包装或审核资料；不凭口感推断"),
        ("商品", "过敏原与交叉接触说明", "以包装及已审核资料为准，不自行推断"),
        ("商品", "营养标签", "按包装标注单位、每份/每100克，避免混用"),
        ("商品", "生产 / 经营资质与许可", "核实主体、商品与经营范围，资料仅限授权用途"),
        ("商品", "生产批次 / 保质期", "记录适用批次与起算条件，不能以旧批次代替"),
        ("商品", "储存方式与食用方法", "与包装要求一致；常温/冷藏等不自行假定"),
        ("商品", "卖点证据与禁用宣称", "口味等主观感受标明；功效、绝对化、产地等须有证据"),
        ("价格", "日常价 / 活动价 / 券后价", "区分商家承担与平台承担优惠及生效区间"),
        ("价格", "达人专属机制与变更权限", "谁可批价、批佣、批固定费，超限交谁批准"),
        ("利润", "单笔售后后净收入口径", "尚未扣本模型费用；勿把已扣佣结算净额再次扣佣"),
        ("利润", "有效计佣销售额与佣金规则", "以合同及后台为准，核对退款、补贴、税费和阶梯规则"),
        ("利润", "平台 / 支付费计提基数", "费率、基数、退款返还等以实际约定为准"),
        ("利润", "货品 / 包材 / 履约 / 售后损失", "完整退款订单的不可回收费用需计入，不重复扣损失"),
        ("履约", "发货地 / 发货时效", "核对承诺、截单时间、预售和节假日例外"),
        ("履约", "可售库存与查询频率", "可售库存、活动库存、预留量分别确认"),
        ("履约", "配送与包邮地区", "偏远地区、运费、无法配送范围"),
        ("履约", "售后口径与升级负责人", "破损、漏发、质量异议、退款责任及举证流程"),
        ("预算", "固定费 / 寄样预算与审批链", "审批额度、财务付款要求、收款主体核验"),
        ("结算", "订单归因窗口", "区分视频、直播、店铺自然、投流，重叠窗口不能直接相加"),
        ("结算", "数据截至与最终结算周期", "支付数据、售后后数据与平台结算数据分开"),
        ("结算", "固定费 / 寄样费用分摊规则", "一合作多视频时只分摊一次，总额必须对得上"),
        ("协作", "商务 / 内容 / 客服 / 供应链联系人", "谁负责，谁审批，异常向谁升级"),
        ("协作", "后台权限与证据访问范围", "按实际授权操作，不分享账号口令或无关个人信息"),
        ("协作", "素材 / 达人内容授权范围", "可否二剪、投流、转载、使用期限与主体"),
        ("目标", "部门 KPI 与考核口径", "待公司确认；不预设触达量、视频量或 GMV 指标"),
    ]
    for row, (category, item, scope) in enumerate(evidence, start=1):
        params.write(row, 0, category, fmt["text"])
        params.write(row, 1, item, fmt["text"])
        params.write(row, 3, scope, fmt["muted"])
        params.write(row, 8, "待核实", fmt["input"])
    params.conditional_format(1, 0, ROWS, 10, {"type": "formula", "criteria": '=AND($H2<>"",$H2<TODAY(),$I2<>"不适用")', "format": fmt["bad"]})

    lead_headers = ["线索ID", "合作ID", "达人UID", "达人名称", "主页链接", "负责人", "线索来源", "内容类目 / 场景", "人群场景匹配1–5", "相关视频播放中位数", "观察样本起日", "观察样本止日", "样本条数", "类目成交证据位置", "退款 / 履约观察", "异常 / 历史合作记录", "合作形式", "固定费报价", "佣金比例", "当前状态", "首次录入日期", "最近跟进日期", "下次跟进日期", "约定发片日期", "拟合作SKU / 版本", "预算审批状态", "审批凭据位置", "沟通结论 / 退出原因", "首次触达日期", "达成合作日期", "商品资料确认日期", "记录检查（自动）"]
    leads = setup("达人线索与合作", lead_headers, {3: 21, 4: 30, 7: 26, 13: 29, 14: 29, 15: 28, 24: 26, 27: 40, 31: 34})
    dates(leads, [10, 11, 20, 21, 22, 23, 28, 29, 30])
    numbers(leads, [9, 12], integer=True)
    numbers(leads, [17]); numbers(leads, [18], percent=True)
    leads.data_validation(1, 8, ROWS, 8, {"validate": "integer", "criteria": "between", "minimum": 1, "maximum": 5, "ignore_blank": True})
    for col, name in [(16, "CollabType"), (19, "LeadStatus"), (25, "ApprovalStatus")]: dropdown(leads, col, name)
    for row in range(2, LAST + 1):
        expression = f'=IF(A{row}="","",IF(COUNTIF($A$2:$A${LAST},A{row})>1,"线索ID重复",IF(AND(B{row}<>"",COUNTIF($B$2:$B${LAST},B{row})>1),"合作ID重复：一合作一行",IF(OR(F{row}="",T{row}="",U{row}=""),"待补负责人/状态/录入日","已建档；继续核对证据"))))'
        formula(leads, row, 31, expression)
    leads.conditional_format(1, 0, ROWS, len(lead_headers)-1, {"type": "formula", "criteria": '=AND($A2<>"",$W2<>"",$W2<TODAY(),$T2<>"复盘完成",$T2<>"终止")', "format": fmt["bad"]})
    leads.write_comment("A1", "线索ID唯一。一个达人可有多次合作；后续合作新建线索/合作行，达人UID保持一致。")
    leads.write_comment("J1", "取明确观察区间内相关视频的中位数，记录样本条数和证据；不只看粉丝量或单条爆款。")

    shipping_headers = ["寄样单ID", "合作ID", "达人UID", "负责人", "SKU / 口味 / 批次", "申请日期", "批准日期", "寄出日期", "快递公司", "运单号", "签收日期", "约定反馈日期", "试用反馈", "约定发片日期", "样品货品成本", "样品包材", "寄样物流费", "寄样总投入（自动）", "当前状态", "下次跟进日期", "异常问题 / 升级人", "审批凭据位置", "收件信息受限存储位置", "签收凭证位置", "商品资料已确认", "记录检查（自动）"]
    shipping = setup("寄样履约", shipping_headers, {4: 30, 12: 35, 20: 40, 22: 32, 25: 32})
    dates(shipping, [5, 6, 7, 10, 11, 13, 19]); numbers(shipping, [14, 15, 16])
    dropdown(shipping, 18, "SampleStatus"); dropdown(shipping, 24, "YesNo")
    for row in range(2, LAST + 1):
        formula(shipping, row, 17, f'=IF(A{row}="","",IF(COUNT(O{row}:Q{row})<3,"待补充",IF(MIN(O{row}:Q{row})<0,"检查输入",SUM(O{row}:Q{row}))))', money=True)
        formula(shipping, row, 25, f'=IF(A{row}="","",IF(COUNTIF($A$2:$A${LAST},A{row})>1,"寄样单ID重复",IF(OR(B{row}="",D{row}="",S{row}=""),"待补合作/负责人/状态",IF(COUNTIF(\'达人线索与合作\'!$B$2:$B${LAST},B{row})<>1,"合作ID未唯一匹配","已关联；核对审批与签收"))))')
    shipping.conditional_format(1, 0, ROWS, len(shipping_headers)-1, {"type": "formula", "criteria": '=AND($A2<>"",$T2<>"",$T2<TODAY(),$S2<>"已完成",$S2<>"已取消")', "format": fmt["bad"]})
    shipping.write_comment("W1", "建议只记录权限受限的公司收件信息表位置，不把手机号和详细地址散落在多个台账里。")

    review_headers = ["归因记录ID", "合作ID", "视频ID", "SKU / 版本", "归因窗口起", "归因窗口止", "数据截至日期", "负责人", "发布日期", "视频链接", "播放量", "商品点击量", "支付订单", "支付GMV", "退款金额", "有效净收入（未扣本表费用）", "有效计佣销售额", "实际佣金", "平台 / 支付费", "有效结算订单", "货品+包材+履约实际成本", "其他变动成本", "固定费分摊", "样品与物流分摊", "贡献估算（自动）", "归因口径证据位置", "复盘结论 / 下轮动作", "下次复盘日期", "当前状态", "是否纳入汇总", "分摊口径说明", "记录检查（自动）", "视频唯一计数（自动）"]
    reviews = setup("视频与订单复盘", review_headers, {9: 30, 15: 27, 20: 28, 25: 32, 26: 40, 30: 36, 31: 42})
    dates(reviews, [4, 5, 6, 8, 27]); numbers(reviews, [10, 11, 12, 19], integer=True)
    numbers(reviews, [13, 14, 15, 16, 17, 18, 20, 21, 22, 23])
    dropdown(reviews, 28, "ReviewStatus"); dropdown(reviews, 29, "YesNo")
    for row in range(2, LAST + 1):
        formula(reviews, row, 24, f'=IF(A{row}="","",IF(COUNT(P{row},R{row}:S{row},U{row}:X{row})<7,"待补充",IF(MIN(P{row},R{row}:S{row},U{row}:X{row})<0,"检查输入",ROUND(P{row}-SUM(R{row}:S{row})-SUM(U{row}:X{row}),4))))', money=True)
        formula(reviews, row, 31,
                f'=IF(A{row}="","",IF(COUNTIF($A$2:$A${LAST},A{row})>1,"归因记录ID重复",IF(OR(B{row}="",C{row}="",D{row}="",H{row}="",COUNT(E{row}:G{row})<3,I{row}="",Z{row}=""),"待补归因/证据/日期",IF(F{row}<E{row},"窗口起止错误",IF(COUNTIF(\'达人线索与合作\'!$B$2:$B${LAST},B{row})<>1,"合作ID未唯一匹配",IF(COUNTIFS($B$2:$B${LAST},B{row},$C$2:$C${LAST},C{row},$D$2:$D${LAST},D{row},$E$2:$E${LAST},E{row},$F$2:$F${LAST},F{row})>1,"窗口重复：应更新原行",IF(AND(AD{row}="是",COUNTIFS($B$2:$B${LAST},B{row},$C$2:$C${LAST},C{row},$D$2:$D${LAST},D{row},$AD$2:$AD${LAST},"是")>1),"重复汇总：同视频SKU仅选一行",IF(AND(AC{row}="已复盘",AD{row}="是",ISNUMBER(Y{row}),AE{row}<>""),"可汇总","观察中/未完成结算核对"))))))))')
        formula(reviews, row, 32, f'=IF(C{row}="","",IF(OR(H{row}="",I{row}=""),"待补充",IF(COUNTIFS($C$2:$C${LAST},C{row},$H$2:$H${LAST},H{row},$I$2:$I${LAST},I{row})<>COUNTIF($C$2:$C${LAST},C{row}),"请统一日期/负责人",IF(COUNTIF($C$2:C{row},C{row})=1,1,0))))')
    reviews.conditional_format(1, 0, ROWS, len(review_headers)-1, {"type": "formula", "criteria": '=AND($A2<>"",$AB2<>"",$AB2<TODAY(),$AC2<>"已复盘",$AC2<>"终止")', "format": fmt["bad"]})
    reviews.conditional_format(1, 31, ROWS, 31, {"type": "text", "criteria": "containing", "value": "重复", "format": fmt["bad"]})
    reviews.write_comment("A1", "一行=合作ID×视频ID×SKU×归因窗口的一份累计快照。同一窗口的次日数据更新原行，不重复追加每日累计值。")
    reviews.write_comment("P1", "填同口径售后后的实际净收入，尚未扣本表R/S/U:X费用。若后台净额已扣佣金或平台费，先还原或改用另一不重复扣费的口径；不要混用。")
    reviews.write_comment("AD1", "同一合作×视频×SKU只允许选一个最终非重复窗口为“是”；观察窗口选“否”。多个独立非重叠订单分组须先核对去重后汇总到唯一行。")
    reviews.write_comment("W1", "同一合作多视频时按事先约定分摊固定费，总分摊额应等于实际固定费；不要每个视频重复扣整笔。样品费同理。")

    quote_headers = ["测算ID", "合作ID", "SKU / 版本", "场景与有效订单口径", "证据 / 参数版本位置", "单笔售后后净收入", "单笔有效计佣销售额", "单笔货品成本", "单笔包材", "单笔履约运费", "单笔其他变动成本", "佣金比例", "平台 / 支付费率", "单笔平台费计费基数", "固定合作费", "样品+包材+寄样物流总投入", "情景有效订单数（可选）", "单笔佣金（自动）", "单笔平台费（自动）", "单笔贡献（自动）", "前置总投入（自动）", "保本有效订单数（自动）", "情景贡献（自动）", "预算审批状态", "负责人", "测算日期", "下次校准日期", "风险备注 / 待核实项"]
    quotes = setup("报价测算", quote_headers, {3: 38, 4: 35, 5: 25, 6: 26, 13: 27, 15: 30, 21: 29, 27: 40})
    numbers(quotes, [5, 6, 7, 8, 9, 10, 13, 14, 15]); numbers(quotes, [11, 12], percent=True); numbers(quotes, [16], integer=True)
    dates(quotes, [25, 26]); dropdown(quotes, 23, "ApprovalStatus")
    for row in range(2, LAST + 1):
        formula(quotes, row, 17, f'=IF(A{row}="","",IF(COUNT(G{row},L{row})<2,"待补充",IF(OR(G{row}<0,L{row}<0,L{row}>1),"检查输入",G{row}*L{row})))', money=True)
        formula(quotes, row, 18, f'=IF(A{row}="","",IF(COUNT(M{row}:N{row})<2,"待补充",IF(OR(M{row}<0,M{row}>1,N{row}<0),"检查输入",M{row}*N{row})))', money=True)
        formula(quotes, row, 19, f'=IF(A{row}="","",IF(COUNT(F{row}:N{row})<9,"待补充",IF(OR(MIN(F{row}:N{row})<0,L{row}>1,M{row}>1),"检查输入",ROUND(F{row}-SUM(H{row}:K{row})-R{row}-S{row},4))))', money=True)
        formula(quotes, row, 20, f'=IF(A{row}="","",IF(COUNT(O{row}:P{row})<2,"待补充",IF(MIN(O{row}:P{row})<0,"检查输入",SUM(O{row}:P{row}))))', money=True)
        formula(quotes, row, 21, f'=IF(A{row}="","",IF(OR(T{row}="检查输入",U{row}="检查输入"),"检查输入",IF(OR(NOT(ISNUMBER(T{row})),NOT(ISNUMBER(U{row}))),"待补充",IF(T{row}<=0,"不可保本",IFERROR(ROUNDUP(U{row}/T{row},0),"检查输入")))))')
        formula(quotes, row, 22, f'=IF(A{row}="","",IF(OR(NOT(ISNUMBER(T{row})),NOT(ISNUMBER(U{row})),NOT(ISNUMBER(Q{row}))),"待补充",IF(OR(Q{row}<0,Q{row}<>INT(Q{row})),"检查输入",ROUND(Q{row}*T{row}-U{row},4))))', money=True)
    quotes.conditional_format(1, 19, ROWS, 19, {"type": "cell", "criteria": "<=", "value": 0, "format": fmt["bad"]})
    quotes.conditional_format(1, 0, ROWS, len(quote_headers)-1, {"type": "formula", "criteria": '=AND($A2<>"",$AA2<>"",$AA2<TODAY())', "format": fmt["bad"]})
    quotes.write_comment("F1", "按每个有效结算订单填写、尚未扣H:K/R:S费用的净收入。已扣佣/平台费的后台净额不能再次扣佣。无法回收的退款/售后损失可摊入K，不与F重复。")
    quotes.write_comment("G1", "计佣销售额未必等于净收入。退款、优惠、补贴和阶梯佣金以合同/后台实际口径确认。")
    quotes.write_comment("N1", "平台/支付费与佣金分别使用自己的计提基数，不默认等于售价。")
    quotes.write_comment("V1", "仅在单笔贡献>0时计算ceil((固定费+全部寄样投入)/单笔贡献)。这是基于假设的静态测算，不是成交预测。")

    weekly_headers = ["周起始日期（手填）", "负责人（空=全部）", "周截止（不含，自动）", "新增线索", "首次触达", "达成合作", "寄出样品单", "签收样品单", "发布视频数（去重）", "本周发片队列已核实净收入", "本周发片队列贡献估算", "本周寄出样品投入", "本周主要问题 / 下周动作", "复盘负责人", "复盘日期", "下次检查日期"]
    weekly = setup("每周漏斗", weekly_headers, {0: 19, 1: 22, 8: 26, 9: 32, 10: 30, 12: 50})
    dates(weekly, [0, 14, 15])
    weekly.set_column(2, 2, 19, fmt["formula"])
    for row in range(2, 54):
        formula(weekly, row, 2, f'=IF(A{row}="","",A{row}+7)')
        weekly.write_formula(row-1, 2, f'=IF(A{row}="","",A{row}+7)', wb.add_format({**base, "bg_color": "#FFF8E8", "num_format": "yyyy-mm-dd", "font_color": "#83632E"}), "")
        def count_event(sheet, idcol, datecol, ownercol):
            common = f"'{sheet}'!${idcol}$2:${idcol}${LAST},\"<>\",'{sheet}'!${datecol}$2:${datecol}${LAST},\">=\"&$A{row},'{sheet}'!${datecol}$2:${datecol}${LAST},\"<\"&$C{row}"
            return f'=IF($A{row}="","",IF($B{row}="",COUNTIFS({common}),COUNTIFS({common},\'{sheet}\'!${ownercol}$2:${ownercol}${LAST},$B{row})))'
        for col, args in [(3, ("达人线索与合作", "A", "U", "F")), (4, ("达人线索与合作", "A", "AC", "F")), (5, ("达人线索与合作", "B", "AD", "F")), (6, ("寄样履约", "A", "H", "D")), (7, ("寄样履约", "A", "K", "D"))]:
            formula(weekly, row, col, count_event(*args))
        # A scalar helper column avoids array-COUNTIF differences between spreadsheet engines.
        s = "'视频与订单复盘'"
        def count_video():
            common = f'{s}!$AG$2:$AG${LAST},{s}!$I$2:$I${LAST},">="&$A{row},{s}!$I$2:$I${LAST},"<"&$C{row}'
            return f'=IF($A{row}="","",IF($B{row}="",SUMIFS({common}),SUMIFS({common},{s}!$H$2:$H${LAST},$B{row})))'
        formula(weekly, row, 8, count_video())
        def sum_checked(sumcol):
            common = f'{s}!${sumcol}$2:${sumcol}${LAST},{s}!$I$2:$I${LAST},">="&$A{row},{s}!$I$2:$I${LAST},"<"&$C{row},{s}!$AF$2:$AF${LAST},"可汇总"'
            return f'=IF($A{row}="","",IF($B{row}="",SUMIFS({common}),SUMIFS({common},{s}!$H$2:$H${LAST},$B{row})))'
        formula(weekly, row, 9, sum_checked("P"), money=True)
        formula(weekly, row, 10, sum_checked("Y"), money=True)
        sample_sum = f'\'寄样履约\'!$R$2:$R${LAST},\'寄样履约\'!$A$2:$A${LAST},"<>",\'寄样履约\'!$H$2:$H${LAST},">="&$A{row},\'寄样履约\'!$H$2:$H${LAST},"<"&$C{row}'
        formula(weekly, row, 11, f'=IF($A{row}="","",IF($B{row}="",SUMIFS({sample_sum}),SUMIFS({sample_sum},\'寄样履约\'!$D$2:$D${LAST},$B{row})))', money=True)
    weekly.write_comment("D1", "本页统计每周发生的里程碑事件，不是同批达人从触达到成交的转化率；不同阶段来自不同批次，不能直接相除。")
    weekly.write_comment("I1", "按视频ID去重，通过视频页AG辅助列计数。同一视频所有观察行的发布日期和负责人必须一致，否则辅助列提示修正且暂不计数。")
    weekly.write_comment("J1", "按本周发布日期划分的队列累计结果，只纳入记录检查为“可汇总”的最终记录。不是本周支付流水、现金到账或店铺总GMV。")
    weekly.write_comment("L1", "未知成本行不会贡献到SUMIFS，必须在周会前核查寄样成本“待补充”，不能把未填当零成本。")

    # A separate illustrative sheet: no sample numbers are placed in working sheets.
    examples = wb.add_worksheet("示例-请勿当真实数据")
    examples.hide_gridlines(2)
    examples.set_tab_color("#D39F54")
    examples.set_column("A:A", 28); examples.set_column("B:B", 24); examples.set_column("C:C", 70)
    examples.merge_range("A1:C2", "仅用于理解公式的虚构示例，不能当作公司价格、佣金、成本或业绩目标。", fmt["warning"])
    example_rows = [
        ("输入项", "纯演示数值", "口径"),
        ("单笔售后后净收入", 20, "未扣下列成本/佣金/平台费；金额均为虚构示例"),
        ("单笔有效计佣销售额", 18, "佣金计提基数单独确认，不默认等于净收入"),
        ("货品 / 包材 / 履约 / 其他", "8 / 1 / 2 / 0.5", "每个有效订单平均承担的成本"),
        ("佣金比例 / 平台费率", "10% / 1%", "仅演示，不是公司政策或平台标准"),
        ("平台费计费基数", 20, "单独输入计费基数"),
        ("固定合作费 / 全部寄样投入", "100 / 20", "前置投入总计120，不能在不同视频重复扣整笔"),
        ("单笔佣金", 1.8, "18×10%"),
        ("单笔平台费", 0.2, "20×1%"),
        ("单笔贡献", 6.5, "20−8−1−2−0.5−1.8−0.2"),
        ("保本有效结算订单数", 19, "ceil(120/6.5)；不是预计一定卖出19单"),
        ("假设有效结算20单的贡献", 10, "20×6.5−120；仅静态场景，不能用支付GMV替代贡献"),
        ("贡献=0或负数", "不可保本", "不做除零，不给负数或无限订单数"),
        ("有未知必填参数", "待补充", "未知成本留空；确认没有该费用时再填0"),
        ("同视频1/3/7天累计窗口", "只选一个最终窗口汇总", "观察窗口标否；同一窗口更新原行，不累计每日快照"),
        ("合作C001 / 视频V001+V002", "固定费分摊合计100", "如按50+50分摊，总计应为100；不能每个视频各填100"),
    ]
    for row, values in enumerate(example_rows, start=3):
        examples.set_row(row, 39)
        for col, val in enumerate(values): examples.write(row, col, val, fmt["header"] if row==3 else fmt["text"])

    instructions = wb.add_worksheet("使用说明")
    instructions.hide_gridlines(2); instructions.set_tab_color("#243E3D")
    instructions.set_column("A:A", 6); instructions.set_column("B:B", 24); instructions.set_column("C:F", 22)
    instructions.merge_range("B2:F3", "达人短视频带货商务 · 精细化台账", fmt["title"])
    instructions.merge_range("B5:F6", "这是可调整的空白工具，不是公司制度、平台规则或业绩承诺。先确认价格利润、审批权限、商品证据和结算归因，再填真实数据。", fmt["warning"])
    instructions.merge_range("B8:F8", "如何开始", fmt["section"])
    guide = [
        ("1 先核资料", "在参数字典确认SKU、配料/营养/净含量、资质、保质储存、发货售后、成本与计费规则；待核实不等于已合规。"),
        ("2 再建档", "达人UID标识账号；线索ID标识一次线索处理；合作ID标识一笔合作，一合作一行。寄样单ID和视频ID分别标识履约与内容。"),
        ("3 状态与跟进", "填写负责人、当前状态、最近动作、下次跟进日期。红底提示日期逾期；按实际结论更新，不能为了清零直接标完成。"),
        ("4 预算前置", "报价测算填写同一SKU/同一口径参数。空白表示未知，0表示已核实不计费。模型金额估算到4位小数，实际佣金/税费/舍入以合同后台为准。"),
        ("5 控制寄样", "审批通过、商品资料确认、收件信息核对后发出。记录真实寄出、签收、反馈、下次跟进与成本。收件详细资料建议放受限公司表，只在此记位置。"),
        ("6 防止重复归因", "复盘一行=合作×视频×SKU×窗口累计快照，同一窗口更新原行。一个合作×视频×SKU只把一个最终窗口标为纳入汇总。记录检查不是完整财务审计。"),
        ("7 费用只扣一次", "复盘P列净收入尚未扣R/S/U:X费用；已扣费的后台结算净额不能再扣相同费用。一合作多视频需分摊固定费和样品费，汇总核对总额。"),
        ("8 周会看事件与队列", "每周漏斗按事件日期统计，不把不同批次直接除成转化率。净收入/贡献是本周发片队列后续已核实结果，不是本周收款。未填成本必须人工追补。"),
        ("9 保存与权限", "工作簿需按公司授权保存。没有自动同步、爬虫、账号登录或后台数据采集。定期备份，分享前移除手机号、地址、成本及未授权数据。"),
        ("10 查看示例", "示例页数字全部虚构，独立于正式表，不参与任何汇总。所有正式工作表默认留空，不预设公司KPI或佣金比例。"),
    ]
    for row, (label, detail) in enumerate(guide, start=9):
        instructions.set_row(row, 65)
        instructions.write(row, 1, label, fmt["text"])
        instructions.merge_range(row, 2, row, 5, detail, fmt["text"])
    instructions.merge_range("B22:F22", "颜色、容量与计算", fmt["section"])
    instructions.merge_range("B23:F24", "浅蓝区域填写；浅黄区域是自动公式，建议不要覆盖。每个台账预留300行，周表预留52周；超过范围需同步延展公式、校验和汇总范围。灰色为说明。可使用表头筛选和冻结列。", fmt["muted"])
    instructions.merge_range("B26:F27", "首次使用请在 Excel / WPS 中允许自动计算并重新计算（通常 F9）。生成工具已校验文件结构、公式引用、空值和保本边界；本工具没有在真实 Excel 中执行公式，不能把结构检查当作实机结算验证。", fmt["muted"])
    instructions.merge_range("B29:F30", "报价测算核心：单笔贡献=售后后净收入−货品−包材−履约−其他变动成本−有效计佣销售额×佣金率−平台费基数×费率；前置投入=固定费+样品及寄样投入；仅贡献>0时保本有效订单=向上取整(前置投入/单笔贡献)。", fmt["text"])
    instructions.freeze_panes(8, 2)
    instructions.set_first_sheet()
    instructions.activate()
    wb.close()
    return output


def validate(output=OUTPUT):
    from openpyxl import load_workbook
    path = Path(output)
    book = load_workbook(path, data_only=False)
    expected = ["使用说明", "参数字典", "达人线索与合作", "寄样履约", "视频与订单复盘", "报价测算", "每周漏斗", "示例-请勿当真实数据"]
    assert set(book.sheetnames) == set(expected)
    assert book.active.title == "使用说明"
    for name in ["达人线索与合作", "寄样履约", "视频与订单复盘", "报价测算"]:
        sheet = book[name]
        assert sheet.freeze_panes == "C2"
        assert sheet.auto_filter.ref and sheet.data_validations.count
        assert sheet["A2"].value is None, f"{name} cannot contain fake operating data"
    quote = book["报价测算"]
    assert 'ROUNDUP(U2/T2,0)' in quote["V2"].value
    assert 'T2<=0,"不可保本"' in quote["V2"].value
    assert 'COUNT(F2:N2)<9' in quote["T2"].value
    assert 'G2*L2' in quote["R2"].value
    assert 'M2*N2' in quote["S2"].value
    assert 'SUM(O2:P2)' in quote["U2"].value
    assert '是否纳入汇总' in book["视频与订单复盘"]["AD1"].value
    assert '$AF$2:$AF$301,"可汇总"' in book["每周漏斗"]["J2"].value
    formula_count = 0
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.data_type == 'f':
                    formula_count += 1
                    assert "#REF!" not in cell.value
                    # Quoted sheet references must point to sheets in this workbook.
                    for referenced in re.findall(r"'([^']+)'!", cell.value): assert referenced in book.sheetnames
    # Cached results are blank until input is entered, rather than fake zeros/GMV.
    cached = load_workbook(path, data_only=True)
    assert cached["报价测算"]["V2"].value in (None, "")
    kwargs = dict(net_income=20, commission_base=18, product_cost=8, packaging=1,
                  shipping=2, other_cost=.5, commission_rate=.1, platform_rate=.01,
                  platform_base=20, fixed_fee=100, sample_budget=20, expected_orders=20)
    result = calculate_quote(**kwargs)
    assert result["unit_contribution"] == Decimal("6.5000")
    assert result["break_even_orders"] == 19
    assert result["scenario_contribution"] == Decimal("10.0000")
    assert calculate_quote(**{**kwargs, "net_income": None})["state"] == "待补充"
    assert calculate_quote(**{**kwargs, "net_income": 13.5})["state"] == "不可保本"
    assert calculate_quote(**{**kwargs, "net_income": 10})["state"] == "不可保本"
    assert calculate_quote(**{**kwargs, "fixed_fee": 0, "sample_budget": 0})["break_even_orders"] == 0
    assert calculate_quote(**{**kwargs, "commission_rate": 1.1})["state"] == "检查输入"
    assert calculate_quote(**{**kwargs, "shipping": -1})["state"] == "检查输入"
    assert calculate_quote(**{**kwargs, "expected_orders": 1.5})["state"] == "检查输入"
    assert calculate_quote(**{**kwargs, "net_income": 23.5})["break_even_orders"] == 12
    report = {"output": str(path), "sheets": book.sheetnames, "formula_cells": formula_count,
              "data_rows_per_ledger": ROWS, "weekly_rows": 52, "reference_boundary_checks": 9,
              "xlsx_bytes": path.stat().st_size,
              "validated": "workbook structure, references, empty caches, independent Decimal model",
              "not_validated": "Excel/WPS real-engine formula recalculation"}
    Path('/tmp/merchant-bd-workbook-validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
    return report


if __name__ == '__main__':
    make_workbook()
    print(json.dumps(validate(), ensure_ascii=False, indent=2))
